"""Чтение Датасеты_хакатон.xlsx (зерно строки = определение параметра).

data_model.md §5.9-5.10:
- лист -> object_type; секции «▌» -> group_name; строка -> parameter_type + object_type_parameter;
- «-» в min/max -> NULL; формулы -> кэш-значение (data_only=True);
- лист «Легенда и использование»: блок «Источники данных» -> object_type.data_source_note,
  блок «Важные допущения» в БД не пишется.
"""
from dataclasses import dataclass
from typing import Any

import openpyxl

from normalizers.strings import norm_text

OBJECT_SHEETS: tuple[str, ...] = ("Склад", "Аэропорт", "Медучреждение")
LEGEND_SHEET = "Легенда и использование"
_SHEET_TITLE_ROWS = 2  # r1 — заголовок листа, r2 — шапка таблицы


@dataclass(frozen=True)
class ParamRow:
    """Одна строка листа параметров (сырые значения ячеек)."""
    name: str
    unit: str | None
    base: Any
    vmin: Any
    vmax: Any
    note: str | None
    group: str | None  # секция «▌ ...» без маркера


def load_objects(path=None) -> tuple[dict[str, list[ParamRow]], dict[str, str]]:
    """Возвращает ({имя_листа: [ParamRow]}, {тип_объекта: источник_данных})."""
    import config

    path = path or config.OBJECTS_XLSX_PATH
    wb = openpyxl.load_workbook(path, data_only=True)  # формулы -> кэш-значения
    missing = [s for s in (*OBJECT_SHEETS, LEGEND_SHEET) if s not in wb.sheetnames]
    if missing:
        raise ValueError(f"в XLSX отсутствуют листы: {missing}; фактические: {wb.sheetnames}")

    sheets: dict[str, list[ParamRow]] = {}
    for sheet_name in OBJECT_SHEETS:
        rows: list[ParamRow] = []
        group: str | None = None
        for r in wb[sheet_name].iter_rows(min_row=_SHEET_TITLE_ROWS + 1, values_only=True):
            name = norm_text(r[0])
            if name is None:
                continue
            if name.startswith("▌"):  # секция -> group_name
                group = name.lstrip("▌").strip()
                continue
            rows.append(
                ParamRow(
                    name=name,
                    unit=norm_text(r[1]),
                    base=r[2],
                    vmin=r[3],
                    vmax=r[4],
                    note=norm_text(r[5]) if len(r) > 5 else None,
                    group=group,
                )
            )
        sheets[sheet_name] = rows

    legend = _read_sources(wb[LEGEND_SHEET])
    return sheets, legend


def _read_sources(ws) -> dict[str, str]:
    """Блок «Источники данных» листа «Легенда и использование» -> {тип: источник}."""
    sources: dict[str, str] = {}
    grabbing = False
    for r in ws.iter_rows(values_only=True):
        a = norm_text(r[0])
        b = norm_text(r[1]) if len(r) > 1 else None
        if a == "Источники данных":
            grabbing = True
            continue
        if a == "Важные допущения":  # дальше — допущения, они в assumptions.md, не в БД
            break
        if grabbing and a and b:
            sources[a] = b
    return sources
